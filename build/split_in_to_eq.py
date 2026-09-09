#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
RuleData xlsx의 IN 연산자 조건을 '=' 단건 조건으로 분해하는 유틸.

업무기준 RuleData 시트(4행 조건명 / 5행 연산자·비교값1·비교값2 / 6행부터 데이터 / 'END!' 이후 각주)에서
연산자가 IN 인 조건을 비교값 1개당 1행으로 쪼갠다. 한 행에 IN 조건이 여러 개면 모든 조합(카티션 곱)을 만든다.
헤더(1~5행)·각주·병합셀·열 너비·지브라 서식·행 높이는 원본 그대로 유지하고, 'no.' 열은 1부터 다시 매긴다.

사용:
  python3 build/split_in_to_eq.py --xlsx sources/설계KEY.xlsx --out sources/설계KEY_IN_to_EQ.xlsx
  python3 build/split_in_to_eq.py --xlsx <입력> --out <출력> --exclude 제품형태,주문용도코드   # 제외 항목 재지정

기본 제외 항목(IN 그대로 유지): 제품형태, 주문용도코드, 고객사코드, 고객사양서번호, 색상코드
원본 xlsx 는 사내 기준정보이므로 저장소에 커밋하지 않는다(`.gitignore: sources/`).
"""

import argparse
import copy
import json
import sys
from itertools import product

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

DEFAULT_EXCLUDE = ["제품형태", "주문용도코드", "고객사코드", "고객사양서번호", "색상코드"]
NAME_ROW, OP_ROW, FIRST_DATA_ROW = 4, 5, 6
END_MARKER = "END!"


def parse_in_values(raw):
    """'A, B,,C' → ['A', 'B', 'C'] (공백 제거·빈 항목 제거·중복 제거, 순서 유지)."""
    parts = [p.strip() for p in str(raw).split(",")]
    return list(dict.fromkeys(p for p in parts if p))


def split_workbook(src, dst, exclude, sheet=None):
    wb = openpyxl.load_workbook(src)
    ws = wb[sheet] if sheet else wb.active
    ncol = ws.max_column

    groups = [(c, ws.cell(NAME_ROW, c).value) for c in range(1, ncol + 1) if ws.cell(OP_ROW, c).value == "연산자"]
    if not groups:
        raise SystemExit("5행에서 '연산자' 헤더를 찾지 못했습니다 — RuleData 양식이 아닙니다.")
    unknown = set(exclude) - {n for _, n in groups}
    if unknown:
        raise SystemExit(f"제외 항목이 4행 조건명에 없습니다: {sorted(unknown)} / 조건명: {[n for _, n in groups]}")
    split_cols = [c for c, n in groups if n not in exclude]

    end_row = next((r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value == END_MARKER), ws.max_row + 1)
    data_rows = range(FIRST_DATA_ROW, end_row)
    footer_rows = range(end_row, ws.max_row + 1)

    records = [([ws.cell(r, c).value for c in range(1, ncol + 1)], ws.row_dimensions[r].height) for r in data_rows]
    footer = [([ws.cell(r, c).value for c in range(1, ncol + 1)],
               [copy.copy(ws.cell(r, c)._style) for c in range(1, ncol + 1)],
               ws.row_dimensions[r].height) for r in footer_rows]
    # 지브라 서식: 출력 짝수행은 원본 6행, 홀수행은 원본 7행 스타일을 따른다(원본이 행 짝·홀로 교차 채움).
    tmpl = {r % 2: [copy.copy(ws.cell(r, c)._style) for c in range(1, ncol + 1)]
            for r in (FIRST_DATA_ROW, FIRST_DATA_ROW + 1)}

    stats = {"source_rows": len(records), "rows_expanded": 0, "in_cells_converted": 0,
             "in_single_value": 0, "in_multi_value": 0, "per_column": {}, "warnings": []}
    out_rows = []
    for k, (vals, height) in enumerate(records):
        src_row = data_rows[k]
        options = []
        for c in split_cols:
            i = c - 1
            op = vals[i]
            if op is None or str(op).strip().upper() != "IN":
                continue
            values = parse_in_values(vals[i + 1]) if vals[i + 1] not in (None, "") else []
            if not values:
                stats["warnings"].append(f"{src_row}행 {get_column_letter(c)}열: IN 비교값이 비어 있어 그대로 둠")
                continue
            if vals[i + 2] not in (None, ""):
                stats["warnings"].append(f"{src_row}행 {get_column_letter(c)}열: IN 에 비교값2={vals[i + 2]!r} 가 있어 비움")
            options.append([(i, v) for v in values])
            stats["in_cells_converted"] += 1
            stats["in_single_value" if len(values) == 1 else "in_multi_value"] += 1
            name = ws.cell(NAME_ROW, c).value
            stats["per_column"][name] = stats["per_column"].get(name, 0) + 1
        if not options:
            out_rows.append((vals, height))
            continue
        stats["rows_expanded"] += 1
        for combo in product(*options):
            nv = list(vals)
            for i, v in combo:
                nv[i], nv[i + 1], nv[i + 2] = "=", v, None
            out_rows.append((nv, height))
    stats["output_rows"] = len(out_rows)

    ws.delete_rows(FIRST_DATA_ROW, ws.max_row - FIRST_DATA_ROW + 1)
    for r in [k for k in ws.row_dimensions if k >= FIRST_DATA_ROW]:
        del ws.row_dimensions[r]

    r = FIRST_DATA_ROW
    for n, (vals, height) in enumerate(out_rows, 1):
        vals[0] = str(n)  # no. 재부여(텍스트)
        sty = tmpl[r % 2]
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            cell.value = vals[c - 1]
            cell._style = copy.copy(sty[c - 1])
        if height is not None:
            ws.row_dimensions[r].height = height
        r += 1
    for vals, styles, height in footer:
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            cell.value = vals[c - 1]
            cell._style = styles[c - 1]
        if height is not None:
            ws.row_dimensions[r].height = height
        r += 1
    wb.save(dst)
    return stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, help="원본 RuleData xlsx 경로")
    ap.add_argument("--out", required=True, help="출력 xlsx 경로")
    ap.add_argument("--exclude", default=",".join(DEFAULT_EXCLUDE),
                    help="IN 을 그대로 둘 조건명(쉼표 구분). 기본: " + ",".join(DEFAULT_EXCLUDE))
    ap.add_argument("--sheet", default=None, help="시트명(기본: 활성 시트)")
    args = ap.parse_args()
    exclude = [x.strip() for x in args.exclude.split(",") if x.strip()]
    stats = split_workbook(args.xlsx, args.out, exclude, args.sheet)
    print(json.dumps(stats, ensure_ascii=False, indent=1))
    print(f"저장: {args.out}")


if __name__ == "__main__":
    main()
